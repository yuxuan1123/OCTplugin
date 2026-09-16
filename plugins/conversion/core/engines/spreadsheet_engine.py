"""
OCTools/core/engines/spreadsheet_engine.py
───────────────────────────────────────────────
表格引擎（核心引擎层）：表格类文件读写（xlsx / csv / json / tsv）

覆盖：
  - 智能读取表格文件 → list[dict]（_read_tabular）
  - list[dict] → 写入目标格式（_write_tabular）
  - xlsx / json / csv 两两互转（xls 统一走 xlsx 逻辑）

设计原则：
  - 每个转换函数签名统一为 (input_path, output_path, log) -> bool
  - 以「list[dict] 记录」为中间表示，读/写分离，便于扩展新格式

旧路径 src/_converters_legacy.py 保留为兼容 shim（re-export 本模块）。
"""
# -*- coding: utf-8 -*-

import os
import csv
import json

from core.utils.file_handler import file_exists


def _read_tabular(path, log):
    """智能读取表格文件 → list[dict]"""
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xls"):
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
        if not rows: return []
        header, data = rows[0], rows[1:]
        return [dict(zip(header, r)) for r in data if any(c is not None for c in r)]
    elif ext == ".csv":
        with open(path, "r", encoding="utf-8-sig") as f:
            return list(csv.DictReader(f))
    elif ext == ".json":
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
        if isinstance(data, list): return data
        if isinstance(data, dict): return [data]
        return []
    elif ext == ".tsv":
        with open(path, "r", encoding="utf-8-sig") as f:
            return list(csv.DictReader(f, delimiter="\t"))
    return []


def _write_tabular(path, records, log):
    """list[dict] → 写入目标格式"""
    if not records:
        log("⚠ 无数据可写入"); return
    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    ext = os.path.splitext(path)[1].lower()
    fields = list(records[0].keys())
    if ext in (".xlsx", ".xls"):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(fields)
        for r in records:
            ws.append([r.get(f, "") for f in fields])
        wb.save(path)
    elif ext == ".csv":
        with open(path, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.DictWriter(f, fieldnames=fields)
            w.writeheader(); w.writerows(records)
    elif ext == ".json":
        with open(path, "w", encoding="utf-8") as f:
            json.dump(records, f, ensure_ascii=False, indent=2)
    elif ext == ".tsv":
        with open(path, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.DictWriter(f, fieldnames=fields, delimiter="\t")
            w.writeheader(); w.writerows(records)


def xlsx_to_json(input_path, output_path, log):
    """XLSX → JSON"""
    if not file_exists(input_path, log): return False
    try:
        log(f"🔄 XLSX → JSON: {input_path}")
        records = _read_tabular(input_path, log)
        _write_tabular(output_path, records, log)
        log(f"✅ 完成 → {output_path}（{len(records)} 条记录）")
        return True
    except Exception as e:
        log(f"❌ {e}"); return False


def xlsx_to_csv(input_path, output_path, log):
    """XLSX → CSV"""
    if not file_exists(input_path, log): return False
    try:
        log(f"🔄 XLSX → CSV: {input_path}")
        records = _read_tabular(input_path, log)
        _write_tabular(output_path, records, log)
        log(f"✅ 完成 → {output_path}")
        return True
    except Exception as e:
        log(f"❌ {e}"); return False


def json_to_xlsx(input_path, output_path, log):
    """JSON → XLSX"""
    if not file_exists(input_path, log): return False
    try:
        log(f"🔄 JSON → XLSX: {input_path}")
        records = _read_tabular(input_path, log)
        _write_tabular(output_path, records, log)
        log(f"✅ 完成 → {output_path}")
        return True
    except Exception as e:
        log(f"❌ {e}"); return False


def json_to_csv(input_path, output_path, log):
    """JSON → CSV"""
    if not file_exists(input_path, log): return False
    try:
        log(f"🔄 JSON → CSV: {input_path}")
        records = _read_tabular(input_path, log)
        _write_tabular(output_path, records, log)
        log(f"✅ 完成 → {output_path}")
        return True
    except Exception as e:
        log(f"❌ {e}"); return False


def csv_to_xlsx(input_path, output_path, log):
    """CSV → XLSX"""
    if not file_exists(input_path, log): return False
    try:
        log(f"🔄 CSV → XLSX: {input_path}")
        records = _read_tabular(input_path, log)
        _write_tabular(output_path, records, log)
        log(f"✅ 完成 → {output_path}")
        return True
    except Exception as e:
        log(f"❌ {e}"); return False


def csv_to_json(input_path, output_path, log):
    """CSV → JSON"""
    if not file_exists(input_path, log): return False
    try:
        log(f"🔄 CSV → JSON: {input_path}")
        records = _read_tabular(input_path, log)
        _write_tabular(output_path, records, log)
        log(f"✅ 完成 → {output_path}")
        return True
    except Exception as e:
        log(f"❌ {e}"); return False


# 便捷别名：xls 统一走 xlsx 逻辑
xls_to_json = xlsx_to_json
xls_to_csv = xlsx_to_csv
json_to_xls = json_to_xlsx
csv_to_xls = csv_to_xlsx